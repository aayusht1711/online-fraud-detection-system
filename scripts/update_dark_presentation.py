import os
import shutil
import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor

def update_text_run(shape, new_text, font_name=None, font_size_pt=None, bold=None, color_rgb=None):
    if not shape.has_text_frame:
        return
    p = shape.text_frame.paragraphs[0]
    # Capture existing formatting if not provided
    if p.runs:
        r_orig = p.runs[0]
        if font_name is None:
            font_name = r_orig.font.name
        if font_size_pt is None and r_orig.font.size:
            font_size_pt = r_orig.font.size.pt
        if bold is None:
            bold = r_orig.font.bold
        if color_rgb is None and r_orig.font.color and hasattr(r_orig.font.color, 'rgb'):
            color_rgb = r_orig.font.color.rgb
            
    p.text = new_text
    if p.runs:
        r = p.runs[0]
        if font_name:
            r.font.name = font_name
        if font_size_pt:
            r.font.size = Pt(font_size_pt)
        if bold is not None:
            r.font.bold = bold
        if color_rgb:
            r.font.color.rgb = color_rgb

def apply_updates(src_path, dst_path):
    print(f"Loading presentation: {src_path}")
    prs = pptx.Presentation(src_path)

    # -------------------------------------------------------------
    # 1. SLIDE 1 (Cover Slide)
    # -------------------------------------------------------------
    s1 = prs.slides[0]
    # Shape 10 (TextBox 11): 'Progress: ~70% Code Complete' -> 'Status: 100% Core Engine Verified'
    update_text_run(
        s1.shapes[10],
        "Status: 100% Core Engine Verified",
        font_name="Inter",
        font_size_pt=9.0,
        bold=False,
        color_rgb=RGBColor(0x34, 0xD3, 0x99)  # Emerald green
    )
    print("  [Slide 1] Updated progress to 100% Core Engine Verified")

    # -------------------------------------------------------------
    # 2. SLIDE 4 (03 / Project Scope)
    # -------------------------------------------------------------
    s4 = prs.slides[3]
    # Shape 28 (TextBox 29): 'Ghost session sandbox drains proxy resources & logs mules.'
    update_text_run(
        s4.shapes[28],
        "2,400ms PoW tarpit drains proxy resources & logs mules.",
        font_name="Inter",
        font_size_pt=8.25,
        bold=False,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    print("  [Slide 4] Updated Honeypot subtext with 2,400ms PoW tarpit")

    # -------------------------------------------------------------
    # 3. SLIDE 6 (05 / Biometric Engineering)
    # -------------------------------------------------------------
    s6 = prs.slides[5]
    # Shape 31 (TextBox 32): '• Bot: H < 0.80 (Straight Vector)' -> '• Bot: H < 0.60 (Collinear Vector)'
    update_text_run(
        s6.shapes[31],
        "• Bot: H < 0.60 (Collinear Vector)",
        font_name="ui-monospace",
        font_size_pt=8.25,
        bold=False,
        color_rgb=RGBColor(0xFB, 0x71, 0x85)  # Rose pink
    )
    # Shape 35 (TextBox 36): '• Paste Bot: T_Flight ≤ 2 ms' -> '• Script Bot: T_Flight ≤ 2.5 ms (0ms Injection)'
    update_text_run(
        s6.shapes[35],
        "• Script Bot: T_Flight \u2264 2.5 ms (0ms Injection)",
        font_name="ui-monospace",
        font_size_pt=8.25,
        bold=False,
        color_rgb=RGBColor(0xFB, 0x71, 0x85)
    )
    print("  [Slide 6] Updated Biometric math thresholds (H < 0.60, Tf <= 2.5ms)")

    # -------------------------------------------------------------
    # 4. SLIDE 7 (06 / Methodology & Timeline)
    # -------------------------------------------------------------
    s7 = prs.slides[6]
    # Shape 25 (TextBox 26): 'Working Prototype: 8 Flask REST APIs, 5 responsive frontend pages, and WebRTC STUN capture.'
    update_text_run(
        s7.shapes[25],
        "Working Prototype: 9 Flask REST APIs, 5 UI portals, Chrome Bot Extension & 100% test pass.",
        font_name="Inter",
        font_size_pt=9.0,
        bold=True,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    print("  [Slide 7] Updated Methodology deliverables to 9 APIs & Chrome Extension")

    # -------------------------------------------------------------
    # 5. SLIDE 8 (07 / System Implementation)
    # -------------------------------------------------------------
    s8 = prs.slides[7]
    # Shape 17 (TextBox 18): '8 REST Endpoints' -> '9 REST Endpoints'
    update_text_run(
        s8.shapes[17],
        "9 REST Endpoints",
        font_name="ui-monospace",
        font_size_pt=9.0,
        bold=False,
        color_rgb=RGBColor(0x93, 0xC5, 0xFD)
    )
    # Shape 19 (TextBox 20): 'Live Demo: /demo (6 Scenarios)' -> 'Live Demo: /demo & Extension (8 Scenarios)'
    update_text_run(
        s8.shapes[19],
        "Live Demo: /demo & Extension (8 Scenarios)",
        font_name="ui-monospace",
        font_size_pt=9.0,
        bold=True,
        color_rgb=RGBColor(0x34, 0xD3, 0x99)
    )
    # Shape 44 (TextBox 45): '• Synthetic ledger balance'
    update_text_run(
        s8.shapes[44],
        "• Synthetic ledger balance (\u20B950k decoy)",
        font_name="ui-monospace",
        font_size_pt=9.0,
        bold=False,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    # Shape 45 (TextBox 46): '• Stalled transfer tarpit'
    update_text_run(
        s8.shapes[45],
        "• 2,400ms PoW tarpit stall challenge",
        font_name="ui-monospace",
        font_size_pt=9.0,
        bold=False,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    # Shape 46 (TextBox 47): '• WebRTC STUN query'
    update_text_run(
        s8.shapes[46],
        "• WebRTC STUN IP capture (103.21.14.88)",
        font_name="ui-monospace",
        font_size_pt=9.0,
        bold=False,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    # Shape 47 (TextBox 48): '• Logs to forensics CSV'
    update_text_run(
        s8.shapes[47],
        "• Traps mule wire: mule_sweep01@okaxis",
        font_name="ui-monospace",
        font_size_pt=9.0,
        bold=False,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    print("  [Slide 8] Updated 9 REST endpoints, Extension demo badge, and Ghost Honeypot specifics")

    # -------------------------------------------------------------
    # 6. SLIDE 9 (08 / Verification & Benchmark)
    # -------------------------------------------------------------
    s9 = prs.slides[8]
    for sh in s9.shapes:
        if sh.has_table:
            t = sh.table
            # Row 7, Col 1: 'Mouse entropy < 0.5, 0ms flight' -> 'Mouse entropy < 0.6, 0ms flight'
            cell_1 = t.cell(7, 1)
            cell_1.text = "Mouse entropy < 0.6, 0ms flight"
            if cell_1.text_frame.paragraphs[0].runs:
                r = cell_1.text_frame.paragraphs[0].runs[0]
                r.font.name = "ui-monospace"
                r.font.size = Pt(8.25)
                r.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
            
            # Row 7, Col 4: 'Tarpit loop & mule capture' -> '2,400ms PoW tarpit & mule capture'
            cell_4 = t.cell(7, 4)
            cell_4.text = "2,400ms PoW tarpit & mule capture"
            if cell_4.text_frame.paragraphs[0].runs:
                r = cell_4.text_frame.paragraphs[0].runs[0]
                r.font.name = "ui-sans-serif"
                r.font.size = Pt(8.25)
                r.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
            break
    print("  [Slide 9] Updated Table Row 7 with H < 0.6 and 2,400ms PoW tarpit")

    # -------------------------------------------------------------
    # 7. SLIDE 10 (09 / Conclusion & Next Phase)
    # -------------------------------------------------------------
    s10 = prs.slides[9]
    # Shape 19 (TextBox 20): 'All 7 Attack Scenarios Verified: ...' -> 'All 8 Attack Scenarios Verified: ...'
    update_text_run(
        s10.shapes[19],
        "All 8 Attack Scenarios Verified: Produced expected deterministic and proportional actions.",
        font_name="Inter",
        font_size_pt=9.0,
        bold=True,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    # Shape 23 (TextBox 24): 'Active Intelligence Extraction: ...'
    update_text_run(
        s10.shapes[23],
        "Active Honeypot Tarpit: Stalls bot automation 2,400ms, extracting origin IPs & mule VPAs.",
        font_name="Inter",
        font_size_pt=9.0,
        bold=True,
        color_rgb=RGBColor(0xCB, 0xD5, 0xE1)
    )
    print("  [Slide 10] Updated conclusion with 8 attack scenarios & 2,400ms Honeypot tarpit")

    prs.save(dst_path)
    print(f"Successfully saved updated dark presentation to: {dst_path}")

if __name__ == "__main__":
    src = r"C:\Users\dell\Downloads\Project_Aegis__Mid-Term_Defense_&_Telemetry_Platform_Original_Backup.pptx"
    dst1 = r"C:\Users\dell\Downloads\Project_Aegis__Mid-Term_Defense_&_Telemetry_Platform.pptx"
    dst2 = r"c:\Users\dell\Desktop\project 2\Project_Aegis__Mid-Term_Defense_&_Telemetry_Platform.pptx"

    apply_updates(src, dst1)
    shutil.copy(dst1, dst2)
    print(f"Copied to project directory: {dst2}")
